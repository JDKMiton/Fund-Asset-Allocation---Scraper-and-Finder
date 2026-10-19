"""Write extracted fund data and look-through results to an Excel workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .allocation import LookThrough
from .config import Template
from .extractor import Extraction

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
PCT = '0.00"%"'


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, min_width: int = 10, max_width: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_width, min(max_width, width + 2))


def _categories(template: Template, extra: list[dict[str, float]]) -> list[str]:
    cats = list(template.category_names)
    for d in extra:
        cats += [k for k in d if k not in cats]
    return cats


def write_workbook(path: str | Path, template: Template, extractions: list[Extraction],
                   results: list[LookThrough]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # 1. Portfolio summary: categories down, portfolios across.
    ws = wb.active
    ws.title = "Look-Through Summary"
    cats = _categories(template, [r.allocation for r in results])
    _header(ws, 1, ["Asset class"] + [r.portfolio for r in results])
    for i, cat in enumerate(cats, start=2):
        ws.cell(row=i, column=1, value=cat)
        for j, r in enumerate(results, start=2):
            ws.cell(row=i, column=j, value=round(r.allocation.get(cat, 0.0), 4)).number_format = PCT
    total_row = len(cats) + 2
    ws.cell(row=total_row, column=1, value="Total").font = TOTAL_FONT
    for j in range(2, len(results) + 2):
        col = get_column_letter(j)
        c = ws.cell(row=total_row, column=j, value=f"=SUM({col}2:{col}{total_row - 1})")
        c.number_format, c.font = PCT, TOTAL_FONT
    warn_row = total_row + 2
    for r in results:
        for w in r.warnings:
            ws.cell(row=warn_row, column=1, value=f"{r.portfolio}: {w}").fill = WARN_FILL
            warn_row += 1
    _autosize(ws)

    # 2. One sheet per portfolio showing each fund's contribution.
    for r in results:
        ws = wb.create_sheet(_sheet_name(f"LT {r.portfolio}", wb))
        cats_r = _categories(template, list(r.contributions.values()))
        _header(ws, 1, ["Fund", "Weight"] + cats_r)
        for i, (fund, contrib) in enumerate(r.contributions.items(), start=2):
            ws.cell(row=i, column=1, value=fund)
            ws.cell(row=i, column=2, value=round(r.weights[fund], 4)).number_format = PCT
            for j, cat in enumerate(cats_r, start=3):
                ws.cell(row=i, column=j, value=round(contrib.get(cat, 0.0), 4)).number_format = PCT
        last = len(r.contributions) + 1
        ws.cell(row=last + 1, column=1, value="Total").font = TOTAL_FONT
        for j in range(2, len(cats_r) + 3):
            col = get_column_letter(j)
            c = ws.cell(row=last + 1, column=j, value=f"=SUM({col}2:{col}{last})")
            c.number_format, c.font = PCT, TOTAL_FONT
        _autosize(ws)

    # 3. Each fund's allocation as extracted from its MDD / fact sheet.
    ws = wb.create_sheet("Fund Allocations")
    cats_f = _categories(template, [e.allocation for e in extractions])
    _header(ws, 1, ["Fund"] + cats_f + ["Total", "As at", "Source", "Notes"])
    for i, e in enumerate(extractions, start=2):
        ws.cell(row=i, column=1, value=e.fund)
        for j, cat in enumerate(cats_f, start=2):
            if cat in e.allocation:
                ws.cell(row=i, column=j, value=e.allocation[cat]).number_format = PCT
        n = len(cats_f) + 2
        ws.cell(row=i, column=n, value=f"=SUM(B{i}:{get_column_letter(n - 1)}{i})").number_format = PCT
        ws.cell(row=i, column=n + 1, value=e.as_at)
        ws.cell(row=i, column=n + 2, value=e.source)
        ws.cell(row=i, column=n + 3, value="; ".join(e.notes))
        if e.notes:
            for j in range(1, n + 4):
                ws.cell(row=i, column=j).fill = WARN_FILL
    _autosize(ws)

    # 4. Audit trail: every label/percentage pair found, and where it was mapped.
    ws = wb.create_sheet("Extracted Lines")
    _header(ws, 1, ["Fund", "Label in document", "Value", "Mapped to"])
    row = 2
    for e in extractions:
        for r in e.raw:
            ws.cell(row=row, column=1, value=e.fund)
            ws.cell(row=row, column=2, value=r.label)
            ws.cell(row=row, column=3, value=r.value).number_format = PCT
            ws.cell(row=row, column=4, value=r.category or "(not mapped)")
            row += 1
    _autosize(ws)

    wb.save(path)
    return path


def _sheet_name(name: str, wb: Workbook) -> str:
    clean = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31]
    candidate, n = clean, 2
    while candidate in wb.sheetnames:
        suffix = f" ({n})"
        candidate, n = clean[:31 - len(suffix)] + suffix, n + 1
    return candidate
