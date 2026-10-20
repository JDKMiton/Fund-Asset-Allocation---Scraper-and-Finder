"""The owner's to-do spreadsheet (folder and tracker fixes): clear the rows marked Done."""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger(__name__)

LIST, DONE = "To do", "Done"


def move_done_rows(path: str | Path, today: date | None = None) -> int:
    """Move the rows whose Status is Done from the "To do" sheet to the "Done" sheet (with the date they were
    moved), so the list only shows open items. Run once a week, so a row marked Done stays visible until the
    next run. Returns the number of rows moved; a missing or locked workbook is logged and skipped."""
    path = Path(path)
    if not path.exists():
        log.warning("to-do list %s not found", path)
        return 0
    wb = load_workbook(path)
    if LIST not in wb.sheetnames:
        log.warning("to-do list %s has no '%s' sheet", path, LIST)
        return 0
    ws = wb[LIST]
    head = [c.value for c in ws[1]]
    if "Status" not in head:
        log.warning("to-do list %s has no Status column", path)
        return 0
    status = head.index("Status")
    done = [r for r in range(2, ws.max_row + 1)
            if str(ws.cell(row=r, column=status + 1).value or "").strip().lower() == "done"]
    if not done:
        return 0
    if DONE not in wb.sheetnames:
        moved = wb.create_sheet(DONE)
        moved.append([*head, "Removed from the list on"])
        for i, c in enumerate(ws[1], start=1):
            moved.cell(row=1, column=i)._style = c._style
            moved.column_dimensions[c.column_letter].width = ws.column_dimensions[c.column_letter].width
    moved = wb[DONE]
    for r in done:
        moved.append([*(ws.cell(row=r, column=c).value for c in range(1, len(head) + 1)),
                      (today or date.today()).isoformat()])
    for r in reversed(done):  # bottom up, so the row numbers still to delete don't move
        ws.delete_rows(r)
    try:
        wb.save(path)
    except PermissionError:  # open in Excel: try again next week
        log.warning("to-do list %s is open, so the %d Done row(s) stay until the next run", path, len(done))
        return 0
    log.info("to-do list: moved %d Done row(s) to the '%s' sheet", len(done), DONE)
    return len(done)
