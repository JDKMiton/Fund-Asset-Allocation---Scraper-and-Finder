import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from examples.make_samples import make_all  # noqa: E402
from fund_allocator.allocation import look_through  # noqa: E402
from fund_allocator.cli import main  # noqa: E402
from fund_allocator.config import Holding, Portfolio, load_template  # noqa: E402
from fund_allocator.extractor import extract_from_pdf, extract_from_text, match_category  # noqa: E402
from fund_allocator.scraper import find_matching_links  # noqa: E402

TEMPLATE = load_template(ROOT / "config" / "template.yaml")


class MatchCategoryTest(unittest.TestCase):
    def test_synonyms(self):
        self.assertEqual(match_category("Local Equity", TEMPLATE), "SA Equity")
        self.assertEqual(match_category("GLOBAL EQUITIES", TEMPLATE), "Foreign Equity")
        self.assertEqual(match_category("Cash & Money Market", TEMPLATE), "SA Cash")
        self.assertEqual(match_category("Offshore cash", TEMPLATE), "Foreign Cash")

    def test_noise_not_mapped(self):
        self.assertIsNone(match_category("Total Expense Ratio (TER)", TEMPLATE))
        self.assertIsNone(match_category("Microsoft", TEMPLATE))


class ExtractTextTest(unittest.TestCase):
    def test_prefers_allocation_section(self):
        text = """Fund X MDD 31 August 2026
        TER 1.5%  1 year return 12.0%
        Asset Allocation
        SA Equity 60.0%   Foreign Equity 30.0%
        Cash 10.0%
        Max offshore 45%"""
        ext = extract_from_text(text, "X", TEMPLATE)
        self.assertEqual(ext.allocation, {"SA Equity": 60.0, "Foreign Equity": 30.0, "SA Cash": 10.0})
        self.assertEqual(ext.as_at, "31 August 2026")
        self.assertEqual(ext.notes, [])

    def test_flags_bad_total(self):
        ext = extract_from_text("Asset allocation\nLocal equity 50%\nForeign equity 20%", "Y", TEMPLATE)
        self.assertTrue(any("70.00%" in n for n in ext.notes))

    def test_nothing_found(self):
        ext = extract_from_text("No numbers here", "Z", TEMPLATE)
        self.assertEqual(ext.allocation, {})
        self.assertTrue(ext.notes)


class PdfTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.balanced, cls.feeder = make_all(Path(cls.tmp.name))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_two_column_pdf(self):
        ext = extract_from_pdf(self.balanced, "Balanced", TEMPLATE)
        self.assertAlmostEqual(ext.total, 100.0)
        self.assertEqual(ext.allocation["SA Equity"], 38.5)
        self.assertEqual(ext.allocation["SA Property"], 4.3)

    def test_comma_decimals(self):
        ext = extract_from_pdf(self.feeder, "Feeder", TEMPLATE)
        self.assertEqual(ext.allocation, {"Foreign Equity": 95.5, "Foreign Cash": 3.0, "SA Cash": 1.5})


class LookThroughTest(unittest.TestCase):
    def test_weighted_sum(self):
        p = Portfolio("P", [Holding("A", 50), Holding("B", 50)])
        funds = {"A": {"SA Equity": 100}, "B": {"SA Equity": 40, "Foreign Equity": 60}}
        r = look_through(p, funds, TEMPLATE)
        self.assertAlmostEqual(r.allocation["SA Equity"], 70)
        self.assertAlmostEqual(r.allocation["Foreign Equity"], 30)
        self.assertEqual(r.warnings, [])

    def test_rescales_weights_and_fund_totals(self):
        p = Portfolio("P", [Holding("A", 30), Holding("B", 30)])
        funds = {"A": {"SA Equity": 99}, "B": {"SA Bonds": 101}}
        r = look_through(p, funds, TEMPLATE)
        self.assertAlmostEqual(r.allocation["SA Equity"], 50)
        self.assertAlmostEqual(sum(r.allocation.values()), 100)
        self.assertTrue(r.warnings)

    def test_missing_fund_is_unallocated(self):
        p = Portfolio("P", [Holding("A", 80), Holding("Missing", 20)])
        r = look_through(p, {"A": {"SA Equity": 100}}, TEMPLATE)
        self.assertAlmostEqual(r.allocation["Unallocated"], 20)


class ScraperLinkTest(unittest.TestCase):
    def test_finds_mdd_link(self):
        html = """<a href="/docs/annual-report.pdf">Annual report</a>
                  <a href="/docs/bal-2026-09.pdf">Minimum Disclosure Document</a>
                  <a href="/docs/fund-factsheet.pdf">Download</a>"""
        links = find_matching_links(html, "https://example.com/fund", r"minimum.?disclosure|fact.?sheet")
        self.assertEqual(links, ["https://example.com/docs/bal-2026-09.pdf",
                                 "https://example.com/docs/fund-factsheet.pdf"])
        # a fund-specific pattern that matches nothing must not fall back to other PDFs on the page
        self.assertEqual(find_matching_links(html, "https://example.com/fund", "abax-absolute-2026-09", fallback=False), [])


class EndToEndTest(unittest.TestCase):
    def test_run_writes_workbook(self):
        make_all()
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            rc = main(["run", "--sources", str(ROOT / "config/sources.yaml"),
                       "--template", str(ROOT / "config/template.yaml"),
                       "--out", str(out), "--download-dir", tmp])
            self.assertEqual(rc, 0)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames[0], "Look-Through Summary")
            self.assertIn("Fund Allocations", wb.sheetnames)
            ws = wb["Look-Through Summary"]
            values = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(2, 14)}
            self.assertAlmostEqual(values["SA Equity"], 23.1)
            self.assertAlmostEqual(values["SA Cash"], 8.625)


class CollectorTest(unittest.TestCase):
    def test_document_month_picks_latest_past_date(self):
        from datetime import date
        from fund_allocator.collector import document_month
        text = "Inception date 1 March 2004\nDate of issue: 30 September 2026\nNext review 2027/01/31"
        self.assertEqual(document_month(text, today=date(2026, 10, 12)), "2026 - 09")
        self.assertIsNone(document_month("no dates", today=date(2026, 10, 12)))

    def test_document_month_ignores_older_period_and_issue_dates(self):
        from datetime import date
        from fund_allocator.collector import document_month
        today = date(2026, 10, 12)
        # TER period end (Dec 2025) before the reporting month
        self.assertEqual(document_month("31 December 2025 August 2026 FUND INFORMATION", today), "2026 - 08")
        # an earlier holdings date in the text, but an explicit "as at end" header
        self.assertEqual(document_month("Top holdings 30 April 2026\nAs at end August 2026", today), "2026 - 08")
        # issue date in the current month is never the reporting month
        self.assertEqual(document_month("Issued 2 October 2026. Fund information on 30 September 2026", today),
                         "2026 - 09")
        # file name is used when the text has no recent date
        self.assertEqual(document_month("Launched 1 May 2001", today, "AGBF - 2026-07.pdf"), "2026 - 07")

    def test_document_month_reads_dotted_dates(self):
        from datetime import date
        from fund_allocator.collector import document_month
        text = "Launch date 23.11.2010\nBenchmark changed 01.08.2022\nAs at 31.08.2026 benchmark is MSCI"
        self.assertEqual(document_month(text, date(2026, 10, 13)), "2026 - 08")

    def test_document_month_decodes_cid_font_text(self):
        from datetime import date
        from fund_allocator.collector import document_month
        # "31 August 2026" encoded as glyph ids offset by 30, as in Marriott's MDDs
        enc = "".join(f"(cid:{ord(c) - 30})" for c in "31 August 2026")
        text = f"Marriott Core Income Fund\n{enc}\nmore text ... September 2026 fee change"
        self.assertEqual(document_month(text, date(2026, 10, 13)), "2026 - 08")

    def test_mdd_facts_from_text(self):
        from fund_allocator.facts import _from_text
        f = _from_text("Inception date: 1 October 2004\nFund size R11.2bn\n"
                       "without taking on undue risk. The Fund's benchmark is the FTSE/JSE All Bond Index.\n")
        self.assertEqual((f["Fund AUM"], f["Inception Date"], f["Benchmark"]),
                         ("R11.2bn", "1 October 2004", "FTSE/JSE All Bond Index"))
        f = _from_text("Fund size: ZAR 19.1bn\nFund inception date: 02.11.87\nBenchmark: CPI + 3% p.a.\n"
                       "Benchmark 8.11 7.07 7.41\n")
        self.assertEqual((f["Fund AUM"], f["Inception Date"], f["Benchmark"]), ("ZAR 19.1bn", "02.11.1987", "CPI + 3% p.a."))
        # long composite benchmarks wrap onto the next line and are kept in full
        f = _from_text("Benchmark: 60% MSCI AC World Net Return, 40%\nFTSE WGBI (60% MSCI Net Return 40% FTSE WGBI\n"
                       "pre 01/01/2011)\nReturns (%) 1yr 3yrs\n")
        self.assertEqual(f["Benchmark"],
                         "60% MSCI AC World Net Return, 40% FTSE WGBI (60% MSCI Net Return 40% FTSE WGBI pre 01/01/2011)")
        long = "From 01/11/2024: " + "; ".join(f"Index {i} ({i}%)" for i in range(1, 30)) + "."
        self.assertEqual(_from_text(f"Benchmark: {long}\n")["Benchmark"], long)
        # fee lines and performance rows are not benchmarks
        self.assertIsNone(_from_text("Benchmark: 1.00% p.a. excl. VAT\nBenchmark 18.15 16.70 14.33")["Benchmark"])

    def test_fund_manager_from_mdd_and_registry(self):
        from fund_allocator.collector import Collected, Fund, write_tracker
        from fund_allocator.config import FundSource
        from fund_allocator.facts import fund_manager
        self.assertEqual(fund_manager("Investment manager: Truffle Asset Management (Pty) Ltd is authorised"),
                         "Truffle Asset Management")
        self.assertEqual(fund_manager("Fund manager Managed by Matrix Fund Managers. An award winning"),
                         "Matrix Fund Managers")
        self.assertIsNone(fund_manager("Investment manager Income distributions Portfolio metrics"))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.xlsx"
            write_tracker(path, "2026 - 08", [Fund("A"), Fund("B")],
                          [Collected("A", "2026 - 08", Path("A.pdf"), facts={"Fund Manager": "Read Ltd"}),
                           Collected("B", "2026 - 08", Path("B.pdf"), facts={"Fund Manager": "Read Ltd"})],
                          sources={"B": FundSource("B", url="x", fund_manager="Veritas Asset Management")})
            ws = load_workbook(path)["Tracker"]
            self.assertEqual([(r[0], r[1]) for r in ws.iter_rows(min_row=2, values_only=True)],
                             [("A", "Read Ltd"), ("B", "Veritas Asset Management")])

    def test_tracker_has_mdd_fact_columns(self):
        from fund_allocator.collector import Collected, Fund, write_tracker
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.xlsx"
            facts = {"Fund AUM": "R1.2bn", "Inception Date": "1 May 2001", "Benchmark": "CPI + 3%"}
            write_tracker(path, "2026 - 08", [Fund("A", 5.0)], [Collected("A", "2026 - 08", Path("A.pdf"), facts=facts)])
            ws = load_workbook(path)["Tracker"]
            self.assertEqual([c.value for c in ws[1]][:8], ["Fund", "Fund Manager", "AUM held", "ISIN", "Fund AUM",
                                                           "Inception Date", "Benchmark", "Status"])
            row = ws[2]
            self.assertEqual([row[4].value, row[5].value.date(), row[6].value, row[7].value],
                             [1200000000, date(2001, 5, 1), "CPI + 3%", "Collected"])
            self.assertEqual((row[4].number_format, row[5].number_format), ('"R" #,##0', "yyyy-mm-dd"))
            # a foreign fund keeps its currency; a unit price read as the fund size is left out (with a note)
            write_tracker(path, "2026 - 08", [Fund("A"), Fund("B")],
                          [Collected("A", "2026 - 08", Path("A.pdf"), facts={"Fund AUM": "US$ 770.18 million",
                                                                             "Inception Date": "January 2005"}),
                           Collected("B", "2026 - 08", Path("B.pdf"), facts={"Fund AUM": "781.34 USD"})])
            ws = load_workbook(path)["Tracker"]
            self.assertEqual((ws["E2"].value, ws["E2"].number_format, ws["F2"].value.date()),
                             (770180000, '"$" #,##0', date(2005, 1, 1)))
            self.assertIn("inception day not printed", ws.cell(row=2, column=ws.max_column).value)
            self.assertIsNone(ws["E3"].value)
            self.assertIn("781.34 USD", ws.cell(row=3, column=ws.max_column).value)

    def test_aum_and_inception_parsing(self):
        from fund_allocator.facts import aum_amount, inception_date
        self.assertEqual(aum_amount("R 9 865 453 307"), (9865453307, "ZAR"))
        self.assertEqual(aum_amount("R361m"), (361000000, "ZAR"))
        self.assertEqual(aum_amount("R 411,622,968.43"), (411622968, "ZAR"))
        self.assertEqual(aum_amount("R 414,360,165 3"), (414360165, "ZAR"))  # footnote digit
        self.assertEqual(aum_amount("GBP 774m"), (774000000, "GBP"))
        self.assertEqual(aum_amount("$4,146.9m"), (4146900000, "USD"))
        self.assertIsNone(aum_amount("8.11bn"))  # no currency
        self.assertIsNone(aum_amount("$18.37"))  # a unit price
        self.assertEqual(inception_date("23November2015"), (date(2015, 11, 23), True))
        self.assertEqual(inception_date("03.04.2000"), (date(2000, 4, 3), True))  # day first
        self.assertEqual(inception_date("2010/08/02"), (date(2010, 8, 2), True))
        self.assertEqual(inception_date("March 15, 2019"), (date(2019, 3, 15), True))
        self.assertEqual(inception_date("Aug 2011"), (date(2011, 8, 1), False))

    def test_expand_url_templates(self):
        from fund_allocator.collector import candidate_urls
        from fund_allocator.config import FundSource
        src = FundSource("X", urls=["https://x/{yyyymm}/f.pdf", "https://x/{yyyy}/{month}/f-{mm}.pdf"])
        self.assertEqual(candidate_urls(src, ["2026 - 08"]),
                         ["https://x/202608/f.pdf", "https://x/2026/august/f-08.pdf"])

    def test_mandg_fact_sheets(self):
        from fund_allocator.scraper import mandg_fact_sheets
        html = (r'{\"name\":\"M\u0026G Balanced\",\"fundId\":\"GLOBALCOMP\",\"files\":[{\"fileType\":\"FactSheet\",'
                r'\"path\":\"https://x/inst.pdf\"}]},{\"name\":\"M\u0026G Balanced Fund\",\"fundId\":\"PRUA\",'
                r'\"files\":[{\"fileType\":\"FactSheet\",\"path\":\"https://x/mandg-balanced-fund.pdf\"},'
                r'{\"fileType\":\"Commentary\",\"path\":\"https://x/c.pdf\"}]}')
        self.assertEqual(mandg_fact_sheets(html)["m&g balanced fund"], "https://x/mandg-balanced-fund.pdf")

    def test_expand_url_keeps_regex_braces(self):
        from fund_allocator.collector import expand_url
        self.assertEqual(expand_url(r"/{Month}-{yyyy}-\d{4}", "2026 - 08"), r"/August-2026-\d{4}")

    def test_expand_url_last_weekday(self):
        from fund_allocator.collector import expand_url
        # 31 October 2026 is a Saturday, 31 August 2026 a Monday
        self.assertEqual(expand_url("{dd}{bdd}", "2026 - 10"), "3130")
        self.assertEqual(expand_url("{dd}{bdd}", "2026 - 08"), "3131")

    def test_find_isins(self):
        from fund_allocator.collector import find_isins
        self.assertEqual(find_isins("ISIN Code ZAE000019808 and IE00B2RGGV79, again ZAE000019808"),
                         ["ZAE000019808", "IE00B2RGGV79"])

    def test_collect_files_by_month_and_tracks_outstanding(self):
        from fund_allocator.collector import Fund, collect
        from fund_allocator.config import FundSource
        from fund_allocator.scraper import Scraper
        with tempfile.TemporaryDirectory() as tmp:
            balanced, feeder = make_all(Path(tmp) / "pdfs")
            funds = [Fund("Balanced", 100.0), Fund("Feeder", 50.0), Fund("No Source", 10.0)]
            sources = {"Balanced": FundSource("Balanced", file=str(balanced)),
                       "Feeder": FundSource("Feeder", file=str(feeder))}
            out = Path(tmp) / "MDDs"
            results = collect(funds, sources, Scraper(Path(tmp) / "dl"), out)
            self.assertTrue((out / "2026 - 09" / "Balanced.pdf").exists())
            ws = load_workbook(out / "2026 - 09" / "!MDD Tracker 2026 - 09.xlsx")["Tracker"]
            status = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=8).value for i in range(2, 5)}
            self.assertEqual(status, {"Balanced": "Collected", "Feeder": "Collected", "No Source": "No source"})
            self.assertEqual(sum(1 for r in results if r.month), 2)

    def test_collect_skips_funds_already_saved_for_all_months(self):
        from fund_allocator.collector import Fund, collect
        from fund_allocator.config import FundSource
        from fund_allocator.scraper import Scraper
        with tempfile.TemporaryDirectory() as tmp:
            balanced, _ = make_all(Path(tmp) / "pdfs")
            out = Path(tmp) / "MDDs"
            (out / "2026 - 08").mkdir(parents=True)
            (out / "2026 - 08" / "Balanced.pdf").write_bytes(balanced.read_bytes())
            (out / "2026 - 08" / "MDD Tracker 2026 - 08.xlsx").write_bytes(b"old")  # replaced by the "!" name
            src = {"Balanced": FundSource("Balanced", file=str(balanced))}
            # September is missing, so the fund is searched and September is filed
            collect([Fund("Balanced")], src, Scraper(Path(tmp) / "dl"), out, months=["2026 - 08", "2026 - 09"])
            self.assertTrue((out / "2026 - 09" / "Balanced.pdf").exists())
            self.assertEqual(sorted(p.name for p in (out / "2026 - 08").glob("*.xlsx")), ["!MDD Tracker 2026 - 08.xlsx"])
            # both months saved: the source isn't touched (a missing file would otherwise be an error)
            src = {"Balanced": FundSource("Balanced", file=str(Path(tmp) / "gone.pdf"), isin="ZAE000000001")}
            results = collect([Fund("Balanced")], src, Scraper(Path(tmp) / "dl"), out,
                              months=["2026 - 08", "2026 - 09"])
            self.assertEqual(sorted(r.month for r in results), ["2026 - 08", "2026 - 09"])
            self.assertTrue(all(r.source.startswith("already saved") for r in results))
            ws = load_workbook(out / "2026 - 09" / "!MDD Tracker 2026 - 09.xlsx")["Tracker"]
            row = [c.value for c in ws[2]]
            self.assertEqual((row[3], row[7]), ("ZAE000000001", "Collected"))

    def test_collect_copies_mdds_to_fund_folders(self):
        from fund_allocator.collector import Fund, collect
        from fund_allocator.config import FundSource
        from fund_allocator.scraper import Scraper
        with tempfile.TemporaryDirectory() as tmp:
            balanced, feeder = make_all(Path(tmp) / "pdfs")
            share = Path(tmp) / "share"
            (share / "Mgr" / "Balanced Folder").mkdir(parents=True)
            (share / "Other" / "Feeder Folder").mkdir(parents=True)
            out = Path(tmp) / "MDDs"
            (out / "2026 - 07").mkdir(parents=True)  # filed by an earlier run: copied too
            (out / "2026 - 07" / "Balanced.pdf").write_bytes(feeder.read_bytes())
            sources = {"Balanced": FundSource("Balanced", file=str(balanced), fund_manager="Mgr",
                                              fund_folder="Balanced Folder", current_name="Balanced Now"),
                       "Feeder": FundSource("Feeder", file=str(feeder), fund_manager="Moved",
                                            fund_folder="Feeder Folder", fund_path="Other/Feeder Folder"),
                       "Lost": FundSource("Lost", file=str(feeder), fund_manager="Mgr", fund_folder="Nowhere")}
            funds = [Fund("Balanced"), Fund("Feeder"), Fund("Lost")]
            issues = []
            collect(funds, sources, Scraper(Path(tmp) / "dl"), out, issues=issues, fund_share=share)
            mdds = share / "Mgr" / "Balanced Folder" / "MDDs"
            self.assertTrue((mdds / "2026" / "2026-09 - Balanced Now.pdf").exists())
            self.assertTrue((mdds / "2026" / "2026-07 - Balanced Now.pdf").exists())
            self.assertTrue((share / "Other" / "Feeder Folder" / "MDDs" / "2026" / "2026-09 - Feeder.pdf").exists())
            self.assertFalse((share / "Mgr" / "Nowhere").exists())  # fund folders are never created
            self.assertIn(("Lost", "Fund folder missing"), {(i.fund, i.kind) for i in issues})
            ws = load_workbook(out / "2026 - 09" / "!MDD Tracker 2026 - 09.xlsx")["Tracker"]
            col = [c.value for c in ws[1]].index("Fund folder copy") + 1
            copy = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=col).value for i in range(2, 5)}
            self.assertEqual(copy["Balanced Now"], str(Path("Mgr/Balanced Folder/MDDs/2026/2026-09 - Balanced Now.pdf")))
            self.assertEqual(copy["Lost"], "fund folder not found")
            # never overwrite: a different file already in the fund folder is kept, the new one saved beside it
            (mdds / "2026" / "2026-07 - Balanced Now.pdf").write_bytes(balanced.read_bytes())
            collect(funds, sources, Scraper(Path(tmp) / "dl"), out, fund_share=share)
            self.assertEqual(len(list((mdds / "2026").glob("2026-07 - Balanced Now (reissued *).pdf"))), 1)
            self.assertEqual((mdds / "2026" / "2026-07 - Balanced Now.pdf").read_bytes(), balanced.read_bytes())

    def test_history_learns_when_an_mdd_appears(self):
        from datetime import date, datetime
        from fund_allocator.collector import History, due_date
        from fund_allocator.config import FundSource
        runs = [(datetime(2026, 8, 3), {"A"}), (datetime(2026, 8, 10), {"A"}),
                (datetime(2026, 9, 7), {"A"}), (datetime(2026, 9, 14), {"A"})]
        first = {("A", "2026 - 07"): datetime(2026, 8, 10, 7), ("A", "2026 - 08"): datetime(2026, 9, 14, 7),
                 ("A", "2026 - 06"): datetime(2026, 8, 3, 7)}  # June: no earlier search, so it doesn't count
        h = History(first, runs)
        self.assertEqual(h.typical_days("A"), (12.0, 2))  # 31 Jul + 10 days, 31 Aug + 14 days
        self.assertEqual(h.collected_on("A", "2026 - 07"), date(2026, 8, 10))
        src = FundSource("A", url="x", lag_days=60)
        self.assertEqual(due_date("2026 - 09", src, h.typical_days("A")), date(2026, 10, 19))  # 30 Sep + 12 + 7
        self.assertEqual(due_date("2026 - 09", src, (12.0, 1)), date(2026, 11, 29))  # too little history: lag_days

    def test_tracker_shows_when_each_mdd_was_collected(self):
        from fund_allocator.collector import Fund, collect
        from fund_allocator.config import FundSource
        from fund_allocator.scraper import Scraper
        with tempfile.TemporaryDirectory() as tmp:
            balanced, _ = make_all(Path(tmp) / "pdfs")
            out = Path(tmp) / "MDDs"
            collect([Fund("Balanced")], {"Balanced": FundSource("Balanced", file=str(balanced))},
                    Scraper(Path(tmp) / "dl"), out)
            ws = load_workbook(out / "2026 - 09" / "!MDD Tracker 2026 - 09.xlsx")["Tracker"]
            head = [c.value for c in ws[1]]
            when = ws.cell(row=2, column=head.index("Collected on") + 1).value
            self.assertEqual(when.date(), date.today())
            self.assertIn("Usually available", head)
            self.assertTrue((out / "_runs.csv").read_text(encoding="utf-8").strip().endswith("Balanced"))

    def test_todo_list_moves_done_rows(self):
        from openpyxl import Workbook
        from fund_allocator.todo import move_done_rows
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "todo.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "To do"
            for row in [("#", "What", "Status"), (1, "rename", "Open"), (2, "delete", "Done"), (3, "move", "done")]:
                ws.append(row)
            wb.save(path)
            self.assertEqual(move_done_rows(path, date(2026, 10, 26)), 2)
            wb = load_workbook(path)
            self.assertEqual([r[1] for r in wb["To do"].iter_rows(min_row=2, values_only=True)], ["rename"])
            self.assertEqual([r[1:] for r in wb["Done"].iter_rows(min_row=2, values_only=True)],
                             [("delete", "Done", "2026-10-26"), ("move", "done", "2026-10-26")])
            self.assertEqual(move_done_rows(path), 0)

    def test_mdd_in_fund_folder_is_not_searched_again(self):
        from fund_allocator.collector import Fund, collect
        from fund_allocator.config import FundSource
        from fund_allocator.scraper import Scraper
        with tempfile.TemporaryDirectory() as tmp:
            balanced, _ = make_all(Path(tmp) / "pdfs")
            share = Path(tmp) / "share"
            kept = share / "Mgr" / "Folder" / "MDDs" / "2026"
            kept.mkdir(parents=True)
            for m in ("08", "09"):
                (kept / f"2026-{m} - Balanced.pdf").write_bytes(balanced.read_bytes())
            out = Path(tmp) / "MDDs"
            gone = str(Path(tmp) / "gone.pdf")  # searching would fail: proves the fund isn't searched
            src = {"Balanced": FundSource("Balanced", file=gone, fund_manager="Mgr", fund_folder="Folder"),
                   "Quarterly": FundSource("Quarterly", file=gone, frequency="quarterly")}
            issues = []
            results = collect([Fund("Balanced")], src, Scraper(Path(tmp) / "dl"), out,
                              months=["2026 - 08", "2026 - 09"], issues=issues, fund_share=share)
            self.assertNotIn("Not found", {i.kind for i in issues})
            self.assertTrue((out / "2026 - 08" / "Balanced.pdf").exists())  # copied back for the tracker
            self.assertEqual(sorted(r.month for r in results), ["2026 - 08", "2026 - 09"])
            # a quarterly fund has nothing due in July and August, so it isn't searched either
            collect([Fund("Quarterly")], src, Scraper(Path(tmp) / "dl"), out,
                    months=["2026 - 07", "2026 - 08"], issues=issues, fund_share=share)
            self.assertNotIn("Not found", {i.kind for i in issues})

    def test_collect_tracks_sources_and_takes_aum_from_workbook(self):
        from openpyxl import Workbook
        make_all()
        with tempfile.TemporaryDirectory() as tmp:
            holdings = Path(tmp) / "aum.xlsx"
            wb = Workbook()
            ws = wb.active
            for row in [("Fund", "AUM"), ("Removed Fund", 5.0), ("Feeder", 7.0), ("Balanced", 100.0)]:
                ws.append(row)
            wb.save(holdings)
            sources = Path(tmp) / "sources.yaml"
            sources.write_text(f"funds:\n- name: Feeder\n  file: {ROOT / 'examples/sample_global_equity_mdd.pdf'}\n"
                               f"- name: Balanced\n  file: {ROOT / 'examples/sample_balanced_mdd.pdf'}\n"
                               f"- name: alpha fund\n  note: no MDD\n")
            out = Path(tmp) / "MDDs"
            main(["collect", "--funds", str(holdings), "--sources", str(sources), "--out", str(out),
                  "--download-dir", tmp, "--month", "2026 - 09"])
            ws = load_workbook(out / "2026 - 09" / "!MDD Tracker 2026 - 09.xlsx")["Tracker"]
            rows = [(r[0], r[2]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
            self.assertEqual(rows, [("alpha fund", None), ("Balanced", 100.0), ("Feeder", 7.0)])


class ChecksTest(unittest.TestCase):
    """Fund-name check, masthead month, statuses, never-overwrite and the manifest."""

    def test_as_pdf(self):
        from fund_allocator.scraper import as_pdf
        self.assertEqual(as_pdf(b"<html>not found</html>" * 100), b"")
        self.assertEqual(as_pdf(b"%PDF-tiny"), b"")  # under 1 KB: an error stub, not an MDD
        wrapped = b"--boundary\r\nContent-Type: application/pdf\r\n\r\n" + b"%PDF-1.4" + b"x" * 2000
        self.assertTrue(as_pdf(wrapped).startswith(b"%PDF-1.4"))

    def test_confirmed_name_needs_the_known_isin_in_the_mdd(self):
        from fund_allocator.collector import confirmed_name
        from fund_allocator.config import FundSource
        body = " text" * 60
        src = FundSource("Sasfin BCI Flexible Income", url="u", isin="ZAE000206397",
                         doc_names=["Otto1890 FR Flexible Income Fund"])
        mdd = "Otto1890 FR Flexible Income Fund (A) ISIN ZAE000206397" + body
        self.assertEqual(confirmed_name(src, mdd), "Otto1890 FR Flexible Income Fund")
        self.assertIsNone(confirmed_name(src, "Otto1890 FR Flexible Income Fund (A)" + body))  # no ISIN printed
        self.assertIsNone(confirmed_name(FundSource("X", url="u", doc_names=["Otto1890 FR Flexible Income Fund"]),
                                         mdd))  # no ISIN on file: the MDD's own ISIN proves nothing
        keep = FundSource("PPS Balanced Passive Fund", url="u", isin="ZAE000206397", keep_name=True,
                          doc_names=["Otto1890 FR Flexible Income Fund"])
        self.assertIsNone(confirmed_name(keep, mdd))

    def test_tracker_shows_the_confirmed_name(self):
        from fund_allocator.collector import Collected, Fund, write_tracker
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.xlsx"
            funds = [Fund("Sasfin BCI Flexible Income"), Fund("Allan Gray Bond")]
            results = [Collected("Sasfin BCI Flexible Income", "2026 - 08", Path("s.pdf"), isin="ZAE000206397",
                                 mdd_name="Otto1890 FR Flexible Income Fund"),
                       Collected("Allan Gray Bond", "2026 - 08", Path("a.pdf"))]
            write_tracker(path, "2026 - 08", funds, results)
            rows = [r for r in load_workbook(path)["Tracker"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual([r[0] for r in rows], ["Allan Gray Bond", "Otto1890 FR Flexible Income Fund"])
            self.assertIn("previously 'Sasfin BCI Flexible Income'", rows[1][-1])
            # an owner-confirmed name is shown without waiting for an ISIN match
            from fund_allocator.config import FundSource
            src = {"Allan Gray Bond": FundSource("Allan Gray Bond", url="u", current_name="Zeta Bond Fund")}
            write_tracker(path, "2026 - 08", funds, results, sources=src)
            rows = [r for r in load_workbook(path)["Tracker"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual([r[0] for r in rows], ["Otto1890 FR Flexible Income Fund", "Zeta Bond Fund"])
            self.assertIn("confirmed by the owner", rows[1][-1])

    def test_names_itself(self):
        from fund_allocator.collector import names_itself
        body = " lorem ipsum" * 30
        self.assertTrue(names_itself("Coronation Balanced Plus Fund | Fund Information" + body,
                                     ["Coronation Balanced Plus Fund"]))
        self.assertFalse(names_itself("Amplify SCI Active Equity Retail Hedge Fund" + body,
                                      ["Amplify SCI Diversified Income RIHF (Terebinth)"]))
        # a renamed fund passes on the name it now prints
        self.assertTrue(names_itself("Merchant West SCI Worldwide Flexible Fund A" + body,
                                     ["Counterpoint SCI Worldwide Flexible Fund", "Merchant West SCI Worldwide Flexible Fund"]))
        self.assertIsNone(names_itself("1 % 1 %", ["Anything"]))  # image-only PDF: can't tell

    def test_masthead_month_ignores_holdings_and_inception_dates(self):
        from datetime import date
        from fund_allocator.collector import document_month, masthead_month
        today = date(2026, 10, 19)
        self.assertEqual(masthead_month("M&G Inflation Plus Fund August 2026 Risk profile", today), "2026 - 08")
        self.assertIsNone(masthead_month("Fund launch date 1 July 2000", today))  # mid-month: not a report date
        text = "Top 10 holdings as at 30 Jun 2026\nAsset allocation as at 31 Aug 2026"
        self.assertEqual(document_month(text, today, masthead="Balanced Fund 31 August 2026"), "2026 - 08")
        # Flagship prints a letter o for the zero of the day
        self.assertEqual(masthead_month("MINIMUM DISCLOSURE DOCUMENT | 3oSEPTEMBER2026", today), "2026 - 09")
        # US fact sheets (iShares) print "as of June 30, 2026"
        self.assertEqual(masthead_month("iShares Silver Trust Fact Sheet as of June 30, 2026", today), "2026 - 06")
        self.assertEqual(document_month("Inception 4/21/2006 Fact Sheet as of June 30, 2026", today), "2026 - 06")

    def test_allocation_stamp(self):
        from fund_allocator.collector import allocation_stamp
        self.assertEqual(allocation_stamp("Asset allocation as at 31 Aug 2026"), "2026 - 08")
        self.assertEqual(allocation_stamp("Asset allocation on 30 June 2026"), "2026 - 06")
        self.assertIsNone(allocation_stamp("no table here"))

    def test_outstanding_status(self):
        from datetime import date
        from fund_allocator.collector import outstanding_status
        from fund_allocator.config import FundSource
        src = FundSource("X", url="https://x/f.pdf")
        self.assertEqual(outstanding_status("2026 - 09", src, date(2026, 10, 19)), "Pending")
        self.assertEqual(outstanding_status("2026 - 08", src, date(2026, 10, 19)), "Late")
        slow = FundSource("X", url="https://x/f.pdf", lag_days=60)
        self.assertEqual(outstanding_status("2026 - 08", slow, date(2026, 10, 19)), "Pending")
        q = FundSource("X", url="https://x/f.pdf", frequency="quarterly")
        self.assertEqual(outstanding_status("2026 - 08", q, date(2026, 10, 19)), "Quarterly")
        self.assertEqual(outstanding_status("2026 - 09", None), "No source")

    def _file(self, tmp, text, content=b"%PDF-1.4 one", name="Alpha Balanced Fund"):
        """Run _file_mdd on a fake PDF whose text is ``text`` (PDF reading is stubbed)."""
        from unittest import mock
        from fund_allocator import collector as C
        from fund_allocator.config import FundSource
        pdf = Path(tmp) / f"dl-{abs(hash(content))}.pdf"
        pdf.write_bytes(content)
        info = {"text": text, "masthead": text[:60], "facts": {}}
        issues, got = [], []
        manifest = C.Manifest(Path(tmp) / "MDDs" / C.MANIFEST)
        with mock.patch.object(C, "_read_pdf", return_value=info):
            filed = C._file_mdd(C.Fund(name), FundSource(name, url="u"), pdf, "https://x/f.pdf",
                                Path(tmp) / "MDDs", ["2026 - 08", "2026 - 09"], got, issues, manifest)
        return filed, issues

    def test_wrong_document_is_rejected_and_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            filed, issues = self._file(tmp, "Beta Equity Fund 31 August 2026 " + "words " * 60)
            self.assertIsNone(filed)
            self.assertEqual([i.kind for i in issues], ["Wrong document"])
            self.assertEqual(len(list((Path(tmp) / "MDDs" / "_rejected").glob("*.pdf"))), 1)
            self.assertFalse((Path(tmp) / "MDDs" / "2026 - 08").exists())

    def test_never_overwrites_a_filed_month(self):
        import csv as _csv
        with tempfile.TemporaryDirectory() as tmp:
            text = "Alpha Balanced Fund 31 August 2026 " + "words " * 60
            filed, _ = self._file(tmp, text, b"%PDF-1.4 first")
            self.assertEqual(filed.month, "2026 - 08")
            again, issues = self._file(tmp, text, b"%PDF-1.4 same text, new bytes")
            self.assertEqual(issues, [])  # same content regenerated: unchanged
            again, issues = self._file(tmp, text.replace("words", "revised"), b"%PDF-1.4 corrected")
            self.assertEqual([i.kind for i in issues], ["Reissued"])
            month = Path(tmp) / "MDDs" / "2026 - 08"
            self.assertEqual((month / "Alpha Balanced Fund.pdf").read_bytes(), b"%PDF-1.4 first")
            self.assertEqual(len(list(month.glob("Alpha Balanced Fund (reissued *).pdf"))), 1)
            with open(Path(tmp) / "MDDs" / "_manifest.csv", newline="") as f:
                results = [r["result"] for r in _csv.DictReader(f)]
            self.assertEqual(results, ["filed", "unchanged", "reissued"])

    def test_stale_copy_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            filed, issues = self._file(tmp, "Alpha Balanced Fund 28 February 2026 " + "words " * 60)
            self.assertIsNone(filed)
            self.assertEqual([i.kind for i in issues], ["Stale copy"])


if __name__ == "__main__":
    unittest.main()
